#!/usr/bin/env python3
"""
Build a narrated introductory lecture (video + slide deck) for one course.

Input (per course folder, e.g. lectures/courses/AI101/):
  course.json                  slide list: layout, title, bullets, cards, diagram, image, narration
  assets/*.png|jpg             optional cinematic images used by slides
  audio/sNN.mp3                narration clip for slide NN (1-based, zero padded)
  audio/sNN_2.mp3, sNN_3.mp3   optional extra narration segments for the same slide

Output (course folder /output, git-ignored):
  <CODE>_<Title>.mp4           1920x1080 narrated lecture with chapter markers
  <CODE>_<Title>.pptx          slide deck (one image per slide) with narration in speaker notes
  <CODE>_<Title>_timing.json   per-slide timing report

Usage:
  python lectures/pipeline/build_lecture.py lectures/courses/AI101
  python lectures/pipeline/build_lecture.py lectures/courses/AI101 --only 1-3 --allow-missing-audio
"""
from __future__ import annotations

import argparse
import json
import math
import re
import shutil
import subprocess
import sys
import textwrap
from pathlib import Path

import imageio_ffmpeg
import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import numpy as np  # noqa: E402
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch  # noqa: E402
from PIL import Image, ImageDraw, ImageFont  # noqa: E402

W, H = 1920, 1080
FFMPEG = imageio_ffmpeg.get_ffmpeg_exe()
AUTHOR = "TRIUNE BIBLICAL UNIVERSITY GLOBAL EXTENSION, USA"
DEGREE = "BACHELOR OF ARTIFICIAL INTELLIGENCE AND ROBOTICS ENGINEERING"

# Palette: deep navy, champagne cream, antique gold
GOLD = (201, 162, 39)
GOLD_LIGHT = (236, 207, 124)
CREAM = (246, 239, 222)
MUTED = (178, 188, 206)
PANEL = (10, 22, 44)
GOLD_HEX, GOLD_L_HEX, CREAM_HEX, NODE_HEX = "#c9a227", "#eccf7c", "#f6efde", "#13233f"

FONT_FILES = {
    "serif_bold": "DejaVuSerif-Bold.ttf",
    "sans": "DejaVuSans.ttf",
    "sans_bold": "DejaVuSans-Bold.ttf",
    "sans_oblique": "DejaVuSans-Oblique.ttf",
}
FONT_DIRS = [Path("/usr/share/fonts/truetype/dejavu"), Path(matplotlib.get_data_path()) / "fonts" / "ttf"]
_font_cache: dict = {}
_measure = ImageDraw.Draw(Image.new("RGB", (10, 10)))


def font(key: str, size: int) -> ImageFont.FreeTypeFont:
    k = (key, size)
    if k not in _font_cache:
        for d in FONT_DIRS:
            p = d / FONT_FILES[key]
            if p.exists():
                _font_cache[k] = ImageFont.truetype(str(p), size)
                break
        else:
            raise FileNotFoundError(FONT_FILES[key])
    return _font_cache[k]


# ----------------------------------------------------------------- text helpers
def wrap(text: str, fnt, max_w: float) -> list[str]:
    lines, cur = [], ""
    for word in text.split():
        trial = f"{cur} {word}".strip()
        if not cur or _measure.textlength(trial, font=fnt) <= max_w:
            cur = trial
        else:
            lines.append(cur)
            cur = word
    if cur:
        lines.append(cur)
    return lines


def spaced_width(text: str, fnt, tracking: float) -> float:
    return sum(_measure.textlength(c, font=fnt) for c in text) + tracking * (len(text) - 1)


def spaced_text(draw: ImageDraw.ImageDraw, xy, text: str, fnt, fill, tracking: float = 0) -> None:
    x, y = xy
    for c in text:
        draw.text((x, y), c, font=fnt, fill=fill)
        x += _measure.textlength(c, font=fnt) + tracking


# ----------------------------------------------------------------- canvas helpers
def cover_crop(img: Image.Image, w: int, h: int) -> Image.Image:
    scale = max(w / img.width, h / img.height)
    nw, nh = math.ceil(img.width * scale), math.ceil(img.height * scale)
    img = img.resize((nw, nh), Image.LANCZOS)
    left, top = (nw - w) // 2, (nh - h) // 2
    return img.crop((left, top, left + w, top + h))


def make_background(image_path: Path | None) -> Image.Image:
    """Cinematic backdrop: navy gradient + gold glow, or a photo with a navy scrim."""
    yy, xx = np.mgrid[0:H, 0:W].astype(np.float32)
    if image_path is not None:
        arr = np.asarray(cover_crop(Image.open(image_path).convert("RGB"), W, H)).astype(np.float32) / 255
    else:
        t = (yy / H)[..., None]
        top, bot = np.array([18, 34, 66]) / 255, np.array([4, 8, 18]) / 255
        arr = top * (1 - t) + bot * t
        d = np.sqrt(((xx - W * 0.86) / W) ** 2 + ((yy - H * 0.05) / H) ** 2)[..., None]
        glow = np.exp(-((d / 0.38) ** 2))
        arr = arr * (1 - 0.30 * glow) + (np.array(GOLD) / 255) * 0.30 * glow
    # navy scrim, stronger on the left where text sits
    alpha = (0.80 - 0.28 * np.linspace(0, 1, W, dtype=np.float32))[None, :, None]
    arr = arr * (1 - alpha) + (np.array([5, 10, 22]) / 255) * alpha
    # vignette
    r = np.sqrt(((xx - W / 2) / (W / 2)) ** 2 + ((yy - H / 2) / (H / 2)) ** 2)[..., None]
    arr = arr * np.clip(1 - 0.30 * r ** 2, 0.6, 1)
    return Image.fromarray((np.clip(arr, 0, 1) * 255).astype(np.uint8))


def overlay(canvas: Image.Image, layer: Image.Image) -> None:
    canvas.paste(Image.alpha_composite(canvas.convert("RGBA"), layer).convert("RGB"))


def glass(canvas: Image.Image, box, radius: int = 26, alpha: int = 165) -> None:
    layer = Image.new("RGBA", canvas.size, (0, 0, 0, 0))
    ImageDraw.Draw(layer).rounded_rectangle(
        box, radius=radius, fill=PANEL + (alpha,), outline=GOLD + (255,), width=2
    )
    overlay(canvas, layer)


def image_panel(canvas: Image.Image, path: Path, box, radius: int = 26) -> None:
    x0, y0, x1, y1 = box
    w, h = x1 - x0, y1 - y0
    tile = cover_crop(Image.open(path).convert("RGB"), w, h).convert("RGBA")
    mask = Image.new("L", (w, h), 0)
    ImageDraw.Draw(mask).rounded_rectangle((0, 0, w - 1, h - 1), radius=radius, fill=255)
    tile.putalpha(mask)
    layer = Image.new("RGBA", canvas.size, (0, 0, 0, 0))
    layer.paste(tile, (x0, y0), tile)
    overlay(canvas, layer)
    ImageDraw.Draw(canvas).rounded_rectangle(box, radius=radius, outline=GOLD, width=3)


def diamond(draw: ImageDraw.ImageDraw, cx: float, cy: float, r: float, fill) -> None:
    draw.polygon([(cx, cy - r), (cx + r, cy), (cx, cy + r), (cx - r, cy)], fill=fill)


def draw_footer(canvas: Image.Image) -> None:
    d = ImageDraw.Draw(canvas)
    d.line([(120, 985), (1800, 985)], fill=GOLD, width=1)
    f = font("sans_bold", 24)
    tw = spaced_width(AUTHOR, f, 4)
    spaced_text(d, ((W - tw) / 2, 1012), AUTHOR, f, GOLD_LIGHT, tracking=4)


def draw_chrome(canvas: Image.Image, course: dict, slide: dict, idx: int, total: int) -> int:
    """Header, title and gold rule for content slides. Returns the y where content may start."""
    d = ImageDraw.Draw(canvas)
    header = f"{course['code']}   |   {course['title'].upper()}"
    spaced_text(d, (120, 50), header, font("sans_bold", 22), GOLD, tracking=3)
    num = f"{idx:02d} / {total:02d}"
    fnum = font("sans_bold", 22)
    spaced_text(d, (1800 - spaced_width(num, fnum, 3), 50), num, fnum, GOLD, tracking=3)
    d.line([(120, 96), (1800, 96)], fill=GOLD, width=1)

    title = slide["title"]
    size = 66
    while size > 44 and _measure.textlength(title, font=font("serif_bold", size)) > 1680:
        size -= 2
    ft = font("serif_bold", size)
    d.text((120, 120), title, font=ft, fill=GOLD_LIGHT)
    bar_y = 120 + int(size * 1.25) + 14
    d.rectangle([120, bar_y, 260, bar_y + 4], fill=GOLD)
    return bar_y + 50


def draw_bullets(canvas: Image.Image, bullets: list[str], box, max_size: int = 40, min_size: int = 22) -> None:
    x0, y0, x1, y1 = box
    text_w = x1 - x0 - 60
    d = ImageDraw.Draw(canvas)
    chosen = None
    for size in range(max_size, min_size - 1, -2):
        f = font("sans", size)
        line_h, gap = int(size * 1.32), int(size * 0.7)
        layouts = [wrap(b, f, text_w) for b in bullets]
        total = sum(len(lines) * line_h for lines in layouts) + gap * (len(bullets) - 1)
        chosen = (size, f, line_h, gap, layouts, total)
        if total <= (y1 - y0):
            break
    size, f, line_h, gap, layouts, total = chosen
    y = y0 + max(0, ((y1 - y0) - total) // 2)
    for lines in layouts:
        diamond(d, x0 + 12, y + size * 0.62, size * 0.24, GOLD)
        for ln in lines:
            d.text((x0 + 44, y), ln, font=f, fill=CREAM)
            y += line_h
        y += gap


def draw_cards(canvas: Image.Image, cards: list[dict], box) -> None:
    x0, y0, x1, y1 = box
    n = len(cards)
    cols = n if n <= 3 else 2
    rows = math.ceil(n / cols)
    gap = 36
    cw = (x1 - x0 - gap * (cols - 1)) // cols
    ch = (y1 - y0 - gap * (rows - 1)) // rows
    d_pos = []
    for i, card in enumerate(cards):
        r, c = divmod(i, cols)
        bx0, by0 = x0 + c * (cw + gap), y0 + r * (ch + gap)
        glass(canvas, (bx0, by0, bx0 + cw, by0 + ch))
        d_pos.append((bx0, by0, card))
    d = ImageDraw.Draw(canvas)
    for bx0, by0, card in d_pos:
        pad = 36
        hf = font("serif_bold", 38 if n <= 3 else 36)
        bf = font("sans", 28 if n <= 3 else 27)
        diamond(d, bx0 + pad + 8, by0 + pad + 26, 9, GOLD)
        y = by0 + pad
        for ln in wrap(card["heading"], hf, cw - 2 * pad - 30):
            d.text((bx0 + pad + 30, y), ln, font=hf, fill=GOLD_LIGHT)
            y += int(hf.size * 1.25)
        y += 18
        for ln in wrap(card["text"], bf, cw - 2 * pad):
            d.text((bx0 + pad, y), ln, font=bf, fill=CREAM)
            y += int(bf.size * 1.35)


# ----------------------------------------------------------------- diagrams (matplotlib)
def _box(ax, cx, cy, bw, bh, text, size=15, fill=NODE_HEX, color=CREAM_HEX, edge=GOLD_HEX, weight="normal"):
    ax.add_patch(
        FancyBboxPatch((cx - bw / 2, cy - bh / 2), bw, bh, boxstyle="round,pad=0,rounding_size=14",
                       fc=fill, ec=edge, lw=2, zorder=2)
    )
    chars = max(8, int(bw * 0.9 / (size * 1.389 * 0.56)))
    ax.text(cx, cy, "\n".join(textwrap.wrap(text, chars)), ha="center", va="center", fontsize=size,
            color=color, zorder=3, weight=weight, linespacing=1.25)


def _arrow(ax, a, b, shrink=0, style="-|>", dashed=False, rad=0.0):
    ax.add_patch(FancyArrowPatch(a, b, arrowstyle=style, mutation_scale=24, lw=2.2, color=GOLD_HEX,
                                 linestyle="--" if dashed else "-", connectionstyle=f"arc3,rad={rad}",
                                 shrinkA=shrink, shrinkB=shrink, zorder=1))


def _cycle(ax, spec, w, h):
    nodes, n = spec["nodes"], len(spec["nodes"])
    bw, bh = min(w * 0.24, 360), max(h * 0.15, 70)
    cx, cy = w / 2, h / 2
    rx, ry = w / 2 - bw / 2 - 30, h / 2 - bh / 2 - 22
    pts = []
    for i in range(n):
        ang = math.radians(90 - i * 360 / n)
        pts.append((cx + rx * math.cos(ang), cy + ry * math.sin(ang)))
    for i in range(n):
        a, b = pts[i], pts[(i + 1) % n]
        dist = math.hypot(b[0] - a[0], b[1] - a[1])
        shrink = min(92, 0.34 * dist / 1.389)
        _arrow(ax, a, b, shrink=shrink, rad=0.18)
    for (x, y), label in zip(pts, nodes):
        _box(ax, x, y, bw, bh, label, size=17, weight="bold")
    if spec.get("center"):
        ax.text(cx, cy, spec["center"], ha="center", va="center", fontsize=22, color=GOLD_L_HEX)


def _timeline(ax, spec, w, h):
    events = spec["nodes"]
    n = len(events)
    y0 = h / 2
    ax.plot([60, w - 60], [y0, y0], color=GOLD_HEX, lw=3, zorder=1)
    xs = np.linspace(120, w - 120, n)
    for i, (x, ev) in enumerate(zip(xs, events)):
        up = i % 2 == 0
        ly = y0 + (h * 0.22 if up else -h * 0.22)
        ax.plot([x, x], [y0, ly], color=GOLD_HEX, lw=1.5, alpha=0.8, zorder=1)
        ax.scatter([x], [y0], s=300, color=GOLD_HEX, edgecolor=CREAM_HEX, lw=2, zorder=3)
        year, _, desc = ev.partition("\n")
        sign = 1 if up else -1
        # year sits next to the stem, the description stacks further away from it
        ax.text(x, ly + sign * 8, year, ha="center", va="bottom" if up else "top",
                fontsize=22, color=GOLD_L_HEX, weight="bold", zorder=3)
        body = "\n".join(textwrap.wrap(desc, 17))
        ax.text(x, ly + sign * 46, body, ha="center", va="bottom" if up else "top",
                fontsize=15, color=CREAM_HEX, linespacing=1.3, zorder=3)


def _tree(ax, spec, w, h):
    root = spec["root"]
    kids = root["children"]
    n = len(kids)
    cx = w / 2
    root_y, root_bh = h - 60, 66
    _box(ax, cx, root_y, 430, root_bh, root["label"], size=21, fill=GOLD_HEX, color=NODE_HEX, weight="bold")
    child_y, cw, ch = h * 0.50, min(230, (w - 200) / n - 14), 80
    xs = np.linspace(120 + cw / 2, w - 120 - cw / 2, n)
    for x, kid in zip(xs, kids):
        mid = (root_y - root_bh / 2 + child_y + ch / 2) / 2
        ax.plot([cx, cx, x, x], [root_y - root_bh / 2, mid, mid, child_y + ch / 2],
                color=GOLD_HEX, lw=2, zorder=1)
        _box(ax, x, child_y, cw, ch, kid["label"], size=15, weight="bold")
        for j, g in enumerate(kid.get("children", [])):
            gy = child_y - 125 - j * 70
            ax.plot([x, x], [child_y - ch / 2, gy + 26], color=GOLD_HEX, lw=2, zorder=1)
            _box(ax, x, gy, cw, 56, g["label"], size=14, edge=GOLD_L_HEX)


def _flow(ax, spec, w, h):
    nodes, n = spec["nodes"], len(spec["nodes"])
    horizontal = w >= h * 1.6
    if horizontal:
        bw = min(300, (w - 120) / n - 56)
        bh = max(110, h * 0.30)
        xs = np.linspace(60 + bw / 2, w - 60 - bw / 2, n)
        ys = [h / 2] * n
    else:
        bw = min(w * 0.56, 400)
        bh = min(96, (h - 80) / n - 44)
        xs = [w / 2] * n
        ys = np.linspace(h - 40 - bh / 2, 40 + bh / 2, n)
    for i, label in enumerate(nodes):
        _box(ax, xs[i], ys[i], bw, bh, label, size=17 if horizontal else 18, weight="bold")
    for i in range(n - 1):
        if horizontal:
            a, b = (xs[i] + bw / 2 + 6, ys[i]), (xs[i + 1] - bw / 2 - 6, ys[i + 1])
        else:
            a, b = (xs[i], ys[i] - bh / 2 - 6), (xs[i + 1], ys[i + 1] + bh / 2 + 6)
        _arrow(ax, a, b)
    if spec.get("loop"):
        src, dst = spec["loop"]
        if horizontal:
            a = (xs[src], ys[src] + bh / 2 + 6)
            b = (xs[dst], ys[dst] + bh / 2 + 6)
        else:
            # return path runs down the right-hand side of the stack
            a = (xs[src] + bw / 2 + 4, ys[src])
            b = (xs[dst] + bw / 2 + 4, ys[dst])
        _arrow(ax, a, b, dashed=True, rad=-0.45 if horizontal else -0.35)
        if spec.get("loop_label"):
            if horizontal:
                lx, ly = (a[0] + b[0]) / 2, (a[1] + b[1]) / 2 + bh * 1.05
                ax.text(lx, ly, spec["loop_label"], ha="center", va="bottom", fontsize=14,
                        color=GOLD_L_HEX, style="italic")
            else:
                lx, ly = a[0] + 22, (a[1] + b[1]) / 2
                ax.text(lx, ly, "\n".join(textwrap.wrap(spec["loop_label"], 16)), ha="left", va="center",
                        fontsize=14, color=GOLD_L_HEX, style="italic", linespacing=1.3)


def _network(ax, spec, w, h):
    layers = spec["layers"]
    labels = spec.get("labels", [])
    xs = np.linspace(w * 0.14, w * 0.86, len(layers))
    pos = []
    for x, n in zip(xs, layers):
        ys = np.linspace(h * 0.80, h * 0.22, n) if n > 1 else [h / 2]
        pos.append([(x, y) for y in ys])
    for li in range(len(layers) - 1):
        for (x1, y1) in pos[li]:
            for (x2, y2) in pos[li + 1]:
                ax.plot([x1, x2], [y1, y2], color=GOLD_HEX, alpha=0.28, lw=1.2, zorder=1)
    for li, layer in enumerate(pos):
        for (x, y) in layer:
            ax.scatter([x], [y], s=900, color=NODE_HEX, edgecolor=GOLD_HEX, lw=2.5, zorder=3)
        if li < len(labels):
            ax.text(xs[li], h * 0.08, labels[li], ha="center", color=CREAM_HEX, fontsize=17)


def make_diagram(spec: dict, w: int, h: int, out: Path) -> Path:
    fig = plt.figure(figsize=(w / 100, h / 100), dpi=100)
    fig.patch.set_alpha(0)
    ax = fig.add_axes([0, 0, 1, 1])
    ax.set_xlim(0, w)
    ax.set_ylim(0, h)
    ax.axis("off")
    kind = spec["type"]
    {"cycle": _cycle, "timeline": _timeline, "tree": _tree, "flow": _flow, "network": _network}[kind](
        ax, spec, w, h)
    fig.savefig(out, transparent=True, dpi=100)
    plt.close(fig)
    return out


# ----------------------------------------------------------------- slide composition
def render_slide(course: dict, slide: dict, idx: int, total: int, cdir: Path, work: Path) -> Image.Image:
    layout = slide["layout"]
    if layout == "cover":
        img_path = cdir / slide["image"] if slide.get("image") else None
        canvas = make_background(img_path).convert("RGB")
        d = ImageDraw.Draw(canvas)
        band = Image.new("RGBA", canvas.size, (0, 0, 0, 0))
        ImageDraw.Draw(band).rectangle([0, 210, W, 900], fill=(4, 9, 20, 150))
        overlay(canvas, band)
        d = ImageDraw.Draw(canvas)
        head_f = font("sans_bold", 26)
        spaced_text(d, ((W - spaced_width(DEGREE, head_f, 5)) / 2, 262), DEGREE, head_f, GOLD_LIGHT, tracking=5)
        size = 110
        while True:
            ft = font("serif_bold", size)
            lines = wrap(slide["title"], ft, 1640)
            if len(lines) <= 2 or size <= 64:
                break
            size -= 4
        y = 340
        for ln in lines:
            tw = _measure.textlength(ln, font=ft)
            d.text(((W - tw) / 2, y), ln, font=ft, fill=CREAM)
            y += int(size * 1.15)
        oy = y + 28
        d.line([(W / 2 - 300, oy), (W / 2 - 50, oy)], fill=GOLD, width=3)
        d.line([(W / 2 + 50, oy), (W / 2 + 300, oy)], fill=GOLD, width=3)
        diamond(d, W / 2, oy, 12, GOLD)
        sf = font("sans", 36)
        tw = _measure.textlength(slide["subtitle"], font=sf)
        d.text(((W - tw) / 2, oy + 44), slide["subtitle"], font=sf, fill=GOLD_LIGHT)
        draw_footer(canvas)
        spaced_text(ImageDraw.Draw(canvas), (120, 50), course["code"], font("sans_bold", 22), GOLD, 3)
        return canvas

    canvas = make_background(None).convert("RGB")
    content_top = draw_chrome(canvas, course, slide, idx, total)
    box_top, box_bottom = 300, 930
    if layout == "bullets":
        if slide.get("image"):
            draw_bullets(canvas, slide["bullets"], (120, box_top, 1000, box_bottom))
            image_panel(canvas, cdir / slide["image"], (1060, box_top, 1800, box_bottom))
        else:
            draw_bullets(canvas, slide["bullets"], (120, box_top, 1800, box_bottom), max_size=46)
    elif layout == "cards":
        draw_cards(canvas, slide["cards"], (160, box_top, 1760, box_bottom))
    elif layout == "diagram":
        glass(canvas, (160, box_top, 1760, 840), radius=28, alpha=140)
        png = make_diagram(slide["diagram"], 1600, 520, work / f"diagram_{idx:02d}.png")
        _paste_rgba(canvas, png, (160, 310))
        if slide.get("caption"):
            cf = font("sans_oblique", 28)
            tw = _measure.textlength(slide["caption"], font=cf)
            ImageDraw.Draw(canvas).text(((W - tw) / 2, 866), slide["caption"], font=cf, fill=MUTED)
    elif layout == "split":
        glass(canvas, (960, box_top, 1800, box_bottom), radius=28, alpha=140)
        png = make_diagram(slide["diagram"], 800, 600, work / f"diagram_{idx:02d}.png")
        draw_bullets(canvas, slide["bullets"], (120, box_top, 900, box_bottom), max_size=38)
        _paste_rgba(canvas, png, (980, 315))
    else:
        raise ValueError(f"unknown layout {layout}")
    draw_footer(canvas)
    return canvas


def _paste_rgba(canvas: Image.Image, png: Path, xy) -> None:
    src = Image.open(png).convert("RGBA")
    layer = Image.new("RGBA", canvas.size, (0, 0, 0, 0))
    layer.paste(src, xy, src)
    overlay(canvas, layer)


# ----------------------------------------------------------------- media helpers
def ff(args: list[str]) -> None:
    subprocess.run([FFMPEG, "-hide_banner", "-loglevel", "error", "-y", *args], check=True)


def probe_duration(path: Path) -> float:
    r = subprocess.run([FFMPEG, "-hide_banner", "-i", str(path)], capture_output=True, text=True)
    m = re.search(r"Duration: (\d+):(\d+):([\d.]+)", r.stderr)
    if not m:
        raise RuntimeError(f"cannot read duration of {path}")
    hh, mm, ss = m.groups()
    return int(hh) * 3600 + int(mm) * 60 + float(ss)


def audio_for_slide(segments: list[Path], out_m4a: Path, pad: float = 0.8) -> None:
    """Concatenate narration segments and append a short breath of silence."""
    inputs: list[str] = []
    for s in segments:
        inputs += ["-i", str(s)]
    n = len(segments)
    fc = "".join(f"[{i}:a]aformat=sample_rates=44100:channel_layouts=stereo[a{i}];" for i in range(n))
    fc += "".join(f"[a{i}]" for i in range(n)) + f"concat=n={n}:v=0:a=1,apad=pad_dur={pad}[out]"
    ff([*inputs, "-filter_complex", fc, "-map", "[out]", "-c:a", "aac", "-b:a", "128k", str(out_m4a)])


def silence_for(seconds: float, out_m4a: Path) -> None:
    ff(["-f", "lavfi", "-i", "anullsrc=r=44100:cl=stereo", "-t", f"{seconds:.2f}", "-c:a", "aac",
        "-b:a", "128k", str(out_m4a)])


def slide_video(jpg: Path, audio: Path, out_mp4: Path) -> None:
    ff(["-loop", "1", "-framerate", "2", "-i", str(jpg), "-i", str(audio),
        "-c:v", "libx264", "-preset", "veryfast", "-crf", "23", "-tune", "stillimage", "-pix_fmt", "yuv420p",
        "-r", "2", "-c:a", "aac", "-b:a", "128k", "-shortest", str(out_mp4)])


def parse_only(spec: str, total: int) -> list[int]:
    if not spec:
        return list(range(1, total + 1))
    out: list[int] = []
    for part in spec.split(","):
        if "-" in part:
            a, b = part.split("-")
            out += range(int(a), int(b) + 1)
        else:
            out.append(int(part))
    return [i for i in out if 1 <= i <= total]


def slugify(text: str) -> str:
    return re.sub(r"[^A-Za-z0-9]+", "_", text).strip("_")


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("course_dir")
    ap.add_argument("--only", default="", help="slide numbers to build, e.g. 1-3,7 (creates a _preview build)")
    ap.add_argument("--allow-missing-audio", action="store_true", help="use estimated silence for missing clips")
    ap.add_argument("--skip-pptx", action="store_true")
    ap.add_argument("--skip-video", action="store_true")
    args = ap.parse_args()

    cdir = Path(args.course_dir).resolve()
    course = json.loads((cdir / "course.json").read_text(encoding="utf-8"))
    slides = course["slides"]
    total = len(slides)
    chosen = parse_only(args.only, total)
    preview = bool(args.only)

    out_dir = cdir / "output"
    work = cdir / "work"
    slides_dir = cdir / "slides"
    for p in (out_dir, work, slides_dir):
        p.mkdir(parents=True, exist_ok=True)

    code_slug = course["code"].replace(".", "")
    base = f"{code_slug}_{slugify(course['title'])}" + ("_preview" if preview else "")

    # 1) render slides
    jpgs: dict[int, Path] = {}
    for i in chosen:
        slide = slides[i - 1]
        img = render_slide(course, slide, i, total, cdir, work)
        jpg = slides_dir / f"s{i:02d}.jpg"
        img.save(jpg, quality=93)
        jpgs[i] = jpg
        print(f"[slide {i:02d}] rendered {slide['layout']:<8} {slide['title']}")

    # 2) narration per slide
    missing = []
    seg_map: dict[int, list[Path]] = {}
    for i in chosen:
        segs = [cdir / "audio" / f"s{i:02d}.mp3"]
        for k, _ in enumerate(slides[i - 1].get("narration_extra", []), start=2):
            segs.append(cdir / "audio" / f"s{i:02d}_{k}.mp3")
        seg_map[i] = segs
        missing += [str(p.relative_to(cdir)) for p in segs if not p.exists()]
    if missing and not args.allow_missing_audio:
        print("Missing narration clips:\n  " + "\n  ".join(missing))
        print("Generate them (or pass --allow-missing-audio for a silent preview).")
        return 2

    # 3) per-slide audio + video
    timing = []
    if not args.skip_video:
        seg_dir = work / "video"
        seg_dir.mkdir(parents=True, exist_ok=True)
        clips = []
        for i in chosen:
            slide = slides[i - 1]
            aud = seg_dir / f"a{i:02d}.m4a"
            real = [p for p in seg_map[i] if p.exists()]
            if real:
                audio_for_slide(real, aud)  # preview builds may use a partial clip set
            else:
                est = len(slide["narration"]) / 860 * 60
                silence_for(est, aud)
            dur = probe_duration(aud)
            vid = seg_dir / f"v{i:02d}.mp4"
            slide_video(jpgs[i], aud, vid)
            clips.append((vid, dur))
            timing.append({"slide": i, "title": slide["title"], "seconds": round(dur, 2),
                           "narration_clips": len(real), "silent_placeholder": not real})
            print(f"[slide {i:02d}] audio+video {dur:6.1f}s")

        # 4) concat + chapters
        listing = seg_dir / "list.txt"
        listing.write_text("".join(f"file '{v.as_posix()}'\n" for v, _ in clips), encoding="utf-8")
        joined = seg_dir / "joined.mp4"
        ff(["-f", "concat", "-safe", "0", "-i", str(listing), "-c", "copy", str(joined)])
        meta = seg_dir / "chapters.txt"
        lines = [";FFMETADATA1"]
        t = 0.0
        for (i, (_, dur)) in zip(chosen, clips):
            start, end = int(t * 1000), int((t + dur) * 1000)
            lines += ["[CHAPTER]", "TIMEBASE=1/1000", f"START={start}", f"END={end}",
                      f"title={i:02d} {slides[i - 1]['title']}"]
            t += dur
        meta.write_text("\n".join(lines) + "\n", encoding="utf-8")
        final = out_dir / f"{base}.mp4"
        ff(["-i", str(joined), "-i", str(meta), "-map_metadata", "1", "-map", "0", "-c", "copy",
            "-movflags", "+faststart", str(final)])
        total_s = t
        print(f"VIDEO  {final.relative_to(cdir)}  {total_s / 60:.2f} min")
        (out_dir / f"{base}_timing.json").write_text(
            json.dumps({"course": course["code"], "total_seconds": round(total_s, 2), "slides": timing}, indent=2),
            encoding="utf-8")

    # 5) PPTX deck with narration in speaker notes
    if not args.skip_pptx:
        from pptx import Presentation
        from pptx.util import Inches

        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
        blank = prs.slide_layouts[6]
        for i in chosen:
            s = prs.slides.add_slide(blank)
            s.shapes.add_picture(str(jpgs[i]), 0, 0, width=prs.slide_width, height=prs.slide_height)
            notes = slides[i - 1]["narration"]
            for extra in slides[i - 1].get("narration_extra", []):
                notes += "\n\n" + extra
            s.notes_slide.notes_text_frame.text = notes
        deck = out_dir / f"{base}.pptx"
        prs.save(str(deck))
        print(f"DECK   {deck.relative_to(cdir)}")

    shutil.rmtree(work, ignore_errors=True)
    return 0


if __name__ == "__main__":
    sys.exit(main())
